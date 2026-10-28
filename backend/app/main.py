import shutil
import re
import sys
import uuid
import logging
from pathlib import Path

if __package__ in (None, ""):
    backend_root = Path(__file__).resolve().parent.parent
    if str(backend_root) not in sys.path:
        sys.path.insert(0, str(backend_root))
    from app.config import settings
    from app.database import Base, engine, get_db
    from app.models import AdminNotice, AdminNotification, Analysis, Candidate, Company, EligibilityRecord, Job, Student, User
    from app.migrations import backup_sqlite_database, run_migrations
    from app.nlp import extract_education, extract_entities
    from app.schemas import AnalysisOut, CandidateOut, AnalysisStatusOut, CompletedResumeOut, CurrentResumeState, JobDescriptionExtractOut, JobCreateIn, JobUpdateIn, JobOut, ShortlistUpdateIn, ShortlistThresholdIn, CandidateIdsIn, SendEmailIn, EmailResponseOut, ShortlistedCandidateOut, CallProcedureIn, FilterCriteriaIn, AuthResponseOut, AuthTokenOut, ChangePasswordIn, LoginIn, RegisterIn, UserOut, RecruiterValidationIn, NoticeIn, NoticeUpdateIn, StudentIn, StudentUpdateIn, CompanyIn, CompanyUpdateIn, EligibilityIn, EligibilityUpdateIn, NotificationIn
    from app.security import create_access_token, hash_password, needs_password_rehash, verify_password
    from app.services import analyse_job
    from app.email_service import EmailService
else:
    from .config import settings
    from .database import Base, engine, get_db
    from .models import AdminNotice, AdminNotification, Analysis, Candidate, Company, EligibilityRecord, Job, Student, User
    from .migrations import backup_sqlite_database, run_migrations
    from .nlp import extract_education, extract_entities
    from .schemas import AnalysisOut, CandidateOut, AnalysisStatusOut, CompletedResumeOut, CurrentResumeState, JobDescriptionExtractOut, JobCreateIn, JobUpdateIn, JobOut, ShortlistUpdateIn, ShortlistThresholdIn, CandidateIdsIn, SendEmailIn, EmailResponseOut, ShortlistedCandidateOut, CallProcedureIn, FilterCriteriaIn, AuthResponseOut, AuthTokenOut, ChangePasswordIn, LoginIn, RegisterIn, UserOut, RecruiterValidationIn, NoticeIn, NoticeUpdateIn, StudentIn, StudentUpdateIn, CompanyIn, CompanyUpdateIn, EligibilityIn, EligibilityUpdateIn, NotificationIn
    from .security import create_access_token, hash_password, needs_password_rehash, verify_password
    from .services import analyse_job
    from .email_service import EmailService

from fastapi import BackgroundTasks, Depends, FastAPI, File, Form, Header, HTTPException, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.openapi.docs import get_swagger_ui_html
from fastapi.openapi.utils import get_openapi
from fastapi.responses import FileResponse, HTMLResponse, StreamingResponse

logger = logging.getLogger(__name__)
from openpyxl import Workbook
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session, selectinload
from datetime import datetime, timezone

backup_sqlite_database(engine)
Base.metadata.create_all(bind=engine)
run_migrations(engine)

app = FastAPI(title=settings.api_title, version=settings.api_version, docs_url=None)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.allowed_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def custom_openapi():
    """Ensure Swagger UI renders the repeated multipart field as file inputs."""
    if app.openapi_schema:
        return app.openapi_schema
    schema = get_openapi(title=app.title, version=app.version, routes=app.routes)
    content = schema["paths"]["/api/v1/analyses"]["post"]["requestBody"]["content"]["multipart/form-data"]
    body_schema = content["schema"]
    if "$ref" in body_schema:
        reference = body_schema["$ref"].rsplit("/", 1)[-1]
        body_schema = schema["components"]["schemas"][reference]
    body_schema["properties"]["resumes"] = {
        "type": "array",
        "items": {"type": "string", "format": "binary"},
        "description": "Select 1–60 PDF resumes; each file must be 5 MB or smaller.",
    }
    content["encoding"] = {"resumes": {"contentType": "application/pdf"}}
    app.openapi_schema = schema
    return app.openapi_schema


app.openapi = custom_openapi


@app.get("/docs", include_in_schema=False)
def swagger_docs():
    """Swagger UI needs a small enhancement for one-picker multi-file selection."""
    response = get_swagger_ui_html(openapi_url=app.openapi_url, title=f"{app.title} - Swagger UI")
    enhancement = """
<script>
  const enableBulkResumePicker = () => {
    document.querySelectorAll('input[type="file"]').forEach((input) => {
      input.multiple = true;
      input.setAttribute('accept', '.pdf,application/pdf');
    });
  };
  new MutationObserver(enableBulkResumePicker).observe(document.body, { childList: true, subtree: true });
  enableBulkResumePicker();
</script>
"""
    return HTMLResponse(response.body.decode("utf-8").replace("</body>", enhancement + "</body>"))


@app.on_event("startup")
def startup() -> None:
    settings.upload_dir.mkdir(parents=True, exist_ok=True)
    backup_sqlite_database(engine)
    Base.metadata.create_all(bind=engine)
    run_migrations(engine)
    # In-process BackgroundTasks cannot survive a service restart. Fail old work
    # explicitly so users can see that a new batch is required.
    with engine.begin() as conn:
        conn.execute(text(
            "UPDATE analyses SET status='failed', current_phase='failed', current_stage='failed', "
            "error='Analysis was interrupted by a server restart. Start a new analysis.', "
            "completed_at=CURRENT_TIMESTAMP WHERE status IN ('queued', 'processing')"
        ))
    seed_default_admin()


@app.get("/health")
def health() -> dict:
    return {
        "status": "ok",
        "max_resumes": settings.max_resumes_per_analysis,
        "max_file_size_bytes": settings.max_file_size_bytes,
        "ocr_available": bool(settings.tesseract_cmd or shutil.which("tesseract")),
    }


VALID_MODULES = {"candidate", "recruiter", "admin_tpo", "super_admin"}
ACTIVE_AUTH_TOKENS: dict[str, tuple[str, datetime]] = {}


def seed_default_admin() -> None:
    if not settings.bootstrap_admin_email or not settings.bootstrap_admin_password:
        return
    db = next(get_db())
    try:
        admin_email = settings.bootstrap_admin_email.lower().strip()
        existing = db.scalar(select(User).where(User.email == admin_email))
        if existing is None:
            db.add(User(
                full_name="Default Admin",
                email=admin_email,
                password_hash=hash_password(settings.bootstrap_admin_password),
                module="admin_tpo",
                role="tpo",
                designation="Admin",
                institution_name="HireSense AI",
                is_approved=True,
                account_status="approved",
            ))
            db.commit()
    finally:
        db.close()


def _make_token(user: User) -> AuthTokenOut:
    token, expires_at = create_access_token()
    ACTIVE_AUTH_TOKENS[token] = (user.email, expires_at)
    return AuthTokenOut(access_token=token, token_type="Bearer")


def _user_to_out(user: User) -> UserOut:
    return UserOut(
        id=user.id,
        full_name=user.full_name,
        email=user.email,
        module=user.module,
        role=user.role,
        phone=user.phone,
        company_name=user.company_name,
        designation=user.designation,
        institution_name=user.institution_name,
        is_approved=user.is_approved,
        created_at=user.created_at,
    )


def _get_authenticated_user(authorization: str | None, db: Session) -> User:
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Authentication required.")
    token = authorization.split(" ", 1)[1].strip()
    token_data = ACTIVE_AUTH_TOKENS.get(token)
    if not token_data:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Your session is invalid or expired.")
    email, expires_at = token_data
    if expires_at <= datetime.utcnow():
        ACTIVE_AUTH_TOKENS.pop(token, None)
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Your session has expired.")
    user = db.scalar(select(User).where(User.email == email))
    if user is None:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "User not found.")
    if not user.is_approved:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This account is not approved for access.")
    return user


def _require_admin_user(authorization: str | None, db: Session) -> User:
    user = _get_authenticated_user(authorization, db)
    if user.module not in {"admin_tpo", "super_admin"}:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Admin access is required.")
    return user


def _require_recruitment_user(authorization: str | None, db: Session) -> User:
    user = _get_authenticated_user(authorization, db)
    if user.module not in {"recruiter", "admin_tpo", "super_admin"}:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Recruiter access is required.")
    return user


def _is_admin(user: User) -> bool:
    return user.module in {"admin_tpo", "super_admin"}


def _get_job_for_user(job_id: str, user: User, db: Session) -> Job:
    job = db.get(Job, job_id)
    if job is None or (not _is_admin(user) and job.owner_user_id != user.id):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Job not found.")
    return job


def _get_analysis_for_user(analysis_id: str, user: User, db: Session, *, include_candidates: bool = False) -> Analysis:
    statement = select(Analysis).where(Analysis.id == analysis_id)
    if include_candidates:
        statement = statement.options(selectinload(Analysis.candidates))
    analysis = db.scalar(statement)
    if analysis is None or (not _is_admin(user) and analysis.owner_user_id != user.id):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Analysis not found.")
    return analysis


def _company_for_job_request(payload_company_id: str | None, user: User, db: Session) -> Company | None:
    if _is_admin(user):
        if payload_company_id is None:
            return None
        company = db.get(Company, payload_company_id)
        if company is None:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Company not found.")
        return company

    if not user.company_name:
        if payload_company_id:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "This account has no approved company association.")
        return None
    company = db.scalar(select(Company).where(Company.name.ilike(user.company_name.strip())))
    if company is None:
        company = Company(name=user.company_name.strip(), status="active")
        db.add(company)
        db.flush()
    if payload_company_id and payload_company_id != company.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "You cannot create a job for another company.")
    return company


def _sync_company_job_count(company_id: str | None, db: Session) -> None:
    if not company_id:
        return
    company = db.get(Company, company_id)
    if company:
        company.job_count = db.query(Job).filter(Job.company_id == company_id, Job.status == "open").count()


def _paginate(query, page: int, page_size: int):
    page = max(page, 1)
    page_size = min(max(page_size, 1), 100)
    return query.offset((page - 1) * page_size).limit(page_size).all()


def _recruiter_out(user: User) -> dict:
    document = None
    if user.verification_document_path:
        document = {
            "original_filename": user.verification_document_name,
            "content_type": user.verification_document_type,
            "file_size": user.verification_document_size or 0,
        }
    return {
        "id": user.id,
        "full_name": user.full_name,
        "email": user.email,
        "phone": user.phone,
        "designation": user.designation,
        "company_name": user.company_name,
        "company_identifier": user.company_identifier,
        "company_email": user.company_email,
        "company_phone": user.company_phone,
        "company_address": user.company_address,
        "company_website": user.company_website,
        "is_verified": user.is_approved,
        "account_status": user.account_status or ("approved" if user.is_approved else "pending"),
        "job_count": 0,
        "created_at": user.created_at,
        "verification_document": document,
    }


def _get_recruiter_or_404(recruiter_id: str, db: Session) -> User:
    user = db.get(User, recruiter_id)
    if user is None or user.module != "recruiter":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Recruiter not found.")
    return user


def _eligibility_out(record: EligibilityRecord, db: Session) -> dict:
    student = db.get(Student, record.candidate_id)
    company = db.get(Company, record.job_id)
    return {
        "id": record.id,
        "candidate_id": record.candidate_id,
        "job_id": record.job_id,
        "student_name": student.full_name if student else "Unknown student",
        "company_name": company.name if company else "Unknown company",
        "job_title": "Open role",
        "status": record.status,
        "notes": record.notes,
        "final_result": record.final_result,
        "breakdown": record.breakdown or [],
        "updated_at": record.updated_at,
        "created_at": record.created_at,
    }


@app.post("/api/auth/login", response_model=AuthResponseOut)
def login_user(payload: LoginIn, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.email == payload.email.lower().strip()))
    if user is None or user.module != payload.module:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This account is not approved for the selected module.")
    if not verify_password(payload.password, user.password_hash):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid email or password.")
    if not user.is_approved:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This account is not approved for access.")
    if needs_password_rehash(user.password_hash):
        user.password_hash = hash_password(payload.password)
        db.commit()
    token = _make_token(user)
    return AuthResponseOut(token=token, user=_user_to_out(user))


@app.post("/api/auth/register", response_model=UserOut, status_code=status.HTTP_201_CREATED)
def register_user(payload: RegisterIn, db: Session = Depends(get_db)):
    normalized_email = payload.email.lower().strip()
    # Administrative accounts are provisioned by a configured bootstrap account;
    # recruiter accounts use the document-verification flow below.
    if payload.module != "candidate":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Use the approved registration flow for this account type.")
    if db.scalar(select(User).where(User.email == normalized_email)):
        raise HTTPException(status.HTTP_409_CONFLICT, "An account with this email already exists.")
    if payload.module == "candidate" and db.scalar(select(Student).where(Student.email == normalized_email)):
        raise HTTPException(status.HTTP_409_CONFLICT, "A student with this email already exists.")

    user = User(
        full_name=payload.full_name,
        email=normalized_email,
        password_hash=hash_password(payload.password),
        module=payload.module,
        role="candidate",
        phone=payload.phone,
        company_name=payload.company_name,
        designation=payload.designation,
        institution_name=payload.institution_name,
        institution_identifier=payload.institution_identifier,
        is_approved=True,
    )
    db.add(user)
    if payload.module == "candidate":
        db.add(Student(
            full_name=payload.full_name,
            email=normalized_email,
            student_id=payload.student_id,
            phone=payload.phone,
        ))
    db.commit()
    db.refresh(user)
    return _user_to_out(user)


@app.post("/api/auth/register/recruiter", response_model=UserOut, status_code=status.HTTP_201_CREATED)
async def register_recruiter(
    full_name: str = Form(...),
    email: str = Form(...),
    password: str = Form(...),
    phone: str | None = Form(None),
    designation: str | None = Form(None),
    company_name: str | None = Form(None),
    company_identifier: str | None = Form(None),
    company_email: str | None = Form(None),
    company_phone: str | None = Form(None),
    company_address: str | None = Form(None),
    company_website: str | None = Form(None),
    verification_document: UploadFile = File(default=None),
    db: Session = Depends(get_db),
):
    normalized_email = email.lower().strip()
    if db.scalar(select(User).where(User.email == normalized_email)):
        raise HTTPException(status.HTTP_409_CONFLICT, "An account with this email already exists.")
    if verification_document is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "A verification document is required.")
    allowed_types = {"application/pdf", "image/jpeg", "image/png"}
    if verification_document.content_type not in allowed_types:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Unsupported verification document type.")
    document_dir = settings.upload_dir / "recruiter-verifications"
    document_dir.mkdir(parents=True, exist_ok=True)
    original_name = Path(verification_document.filename or "verification-document").name
    destination = document_dir / f"{uuid.uuid4()}-{original_name}"
    size = 0
    with destination.open("wb") as output:
        while chunk := await verification_document.read(1024 * 1024):
            size += len(chunk)
            if size > 10 * 1024 * 1024:
                output.close()
                destination.unlink(missing_ok=True)
                raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "The verification document must be 10 MB or smaller.")
            output.write(chunk)
    user = User(
        full_name=full_name,
        email=normalized_email,
        password_hash=hash_password(password),
        module="recruiter",
        role="recruiter",
        phone=phone,
        company_name=company_name,
        company_identifier=company_identifier,
        company_email=company_email,
        company_phone=company_phone,
        company_address=company_address,
        company_website=company_website,
        designation=designation,
        is_approved=False,
        account_status="pending",
        verification_document_path=str(destination),
        verification_document_name=original_name,
        verification_document_type=verification_document.content_type,
        verification_document_size=size,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return _user_to_out(user)


@app.get("/api/auth/me", response_model=UserOut)
def get_current_user_info(authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    user = _get_authenticated_user(authorization, db)
    return _user_to_out(user)


@app.post("/api/auth/change-password")
def change_password(payload: ChangePasswordIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    user = _get_authenticated_user(authorization, db)
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Your current password is incorrect.")
    user.password_hash = hash_password(payload.new_password)
    db.commit()
    return {"message": "Password updated successfully."}


@app.get("/api/admin/dashboard")
def admin_dashboard(authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    students = db.query(Student).count()
    recruiters = db.query(User).filter(User.module == "recruiter").count()
    verified_recruiters = db.query(User).filter(User.module == "recruiter", User.is_approved == True).count()
    pending_recruiters = db.query(User).filter(User.module == "recruiter", User.account_status == "pending").count()
    companies = db.query(Company).count()
    jobs = sum(company.job_count or 0 for company in db.scalars(select(Company)).all())
    eligibility_assessments = db.query(EligibilityRecord).count()
    eligible_records = db.query(EligibilityRecord).filter(EligibilityRecord.status == "eligible").count()
    applications = db.query(Candidate).count()
    rate = round((eligible_records / eligibility_assessments) * 100) if eligibility_assessments else 0
    recent_recruiters = [
        {"id": user.id, "company": user.company_name or "Unassigned", "contact": user.full_name, "jobs": 0, "status": user.account_status or ("verified" if user.is_approved else "pending")}
        for user in db.scalars(select(User).where(User.module == "recruiter").order_by(User.created_at.desc()).limit(5)).all()
    ]
    distribution = [
        {"name": "Eligible", "value": eligible_records},
        {"name": "Review", "value": db.query(EligibilityRecord).filter(EligibilityRecord.status == "review").count()},
        {"name": "Not eligible", "value": db.query(EligibilityRecord).filter(EligibilityRecord.status == "not eligible").count()},
    ]
    return {
        "students": students,
        "applications": applications,
        "verified_recruiters": verified_recruiters,
        "recruiters": recruiters,
        "pending_recruiters": pending_recruiters,
        "companies": companies,
        "jobs": jobs,
        "eligible_records": eligible_records,
        "eligibility_assessments": eligibility_assessments,
        "eligibility_rate": rate,
        "notices_expiring_this_week": db.query(AdminNotice).filter(AdminNotice.expires_at != None).count(),
        "eligibility_reviews": db.query(EligibilityRecord).filter(EligibilityRecord.status == "review").count(),
        "recent_recruiters": recent_recruiters,
        "placement_distribution": distribution,
        "application_trend": [
            {"month": "Jan", "applications": 0},
            {"month": "Feb", "applications": 0},
            {"month": "Mar", "applications": 0},
            {"month": "Apr", "applications": 0},
            {"month": "May", "applications": 0},
            {"month": "Jun", "applications": applications},
        ],
        "application_status_distribution": [
            {"name": "Shortlisted", "value": db.query(Candidate).filter(Candidate.is_shortlisted == True).count()},
            {"name": "In Review", "value": max(applications - db.query(Candidate).filter(Candidate.is_shortlisted == True).count(), 0)},
        ],
    }


@app.get("/api/admin/recruiters")
def admin_recruiters(page: int = 1, page_size: int = 10, search: str = "", status_filter: str = "", authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    query = db.query(User).filter(User.module == "recruiter")
    if search:
        like = f"%{search.lower()}%"
        query = query.filter((User.full_name.ilike(like)) | (User.email.ilike(like)) | (User.company_name.ilike(like)))
    if status_filter:
        if status_filter == "verified":
            query = query.filter(User.is_approved == True)
        else:
            query = query.filter(User.account_status == status_filter)
    total = query.count()
    users = query.order_by(User.created_at.desc()).offset((page - 1) * page_size).limit(page_size).all()
    return {"items": [_recruiter_out(user) for user in users], "total": total, "pending_total": db.query(User).filter(User.module == "recruiter", User.account_status == "pending").count(), "verified_total": db.query(User).filter(User.module == "recruiter", User.is_approved == True).count()}


@app.get("/api/admin/recruiters/{recruiter_id}")
def admin_recruiter(recruiter_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    user = _get_recruiter_or_404(recruiter_id, db)
    return _recruiter_out(user)


@app.get("/api/admin/recruiters/{recruiter_id}/verification-document")
def recruiter_verification_document(recruiter_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    user = _get_recruiter_or_404(recruiter_id, db)
    if not user.verification_document_path or not Path(user.verification_document_path).exists():
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Verification document not found.")
    return FileResponse(user.verification_document_path, media_type=user.verification_document_type or "application/octet-stream", filename=user.verification_document_name or "verification-document")


@app.patch("/api/admin/recruiters/{recruiter_id}/validation")
def update_recruiter_validation(recruiter_id: str, body: RecruiterValidationIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    user = _get_recruiter_or_404(recruiter_id, db)
    is_verified = body.is_verified
    user.is_approved = is_verified
    user.account_status = "approved" if is_verified else "rejected"
    db.commit(); db.refresh(user)
    return {**_recruiter_out(user), "message": "Recruiter validation updated."}


@app.get("/api/admin/notices")
def admin_notices(page: int = 1, page_size: int = 10, search: str = "", status_filter: str = "", authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    query = db.query(AdminNotice)
    if search:
        like = f"%{search}%"
        query = query.filter((AdminNotice.title.ilike(like)) | (AdminNotice.content.ilike(like)))
    if status_filter:
        query = query.filter(AdminNotice.status == status_filter)
    total = query.count()
    return {"items": _paginate(query.order_by(AdminNotice.created_at.desc()), page, page_size), "total": total}


@app.get("/api/admin/notices/{notice_id}")
def admin_notice(notice_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    notice = db.get(AdminNotice, notice_id)
    if notice is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Notice not found.")
    return notice


@app.post("/api/admin/notices")
def create_notice(payload: NoticeIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    notice = AdminNotice(**payload.model_dump())
    db.add(notice); db.commit(); db.refresh(notice)
    return notice


@app.patch("/api/admin/notices/{notice_id}")
def update_notice(notice_id: str, payload: NoticeUpdateIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    notice = db.get(AdminNotice, notice_id)
    if notice is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Notice not found.")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(notice, key, value)
    db.commit(); db.refresh(notice)
    return notice


@app.patch("/api/admin/notices/{notice_id}/publish")
def publish_notice(notice_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    notice = db.get(AdminNotice, notice_id)
    if notice is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Notice not found.")
    notice.status = "published"; notice.published_at = datetime.utcnow()
    db.commit(); db.refresh(notice)
    return notice


@app.delete("/api/admin/notices/{notice_id}")
def delete_notice(notice_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    notice = db.get(AdminNotice, notice_id)
    if notice is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Notice not found.")
    db.delete(notice); db.commit()
    return {"id": notice_id, "deleted": True}


@app.get("/api/admin/students")
def admin_students(page: int = 1, page_size: int = 10, search: str = "", department: str = "", eligibility_status: str = "", application_status: str = "", authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    query = db.query(Student)
    if search:
        like = f"%{search}%"
        query = query.filter((Student.full_name.ilike(like)) | (Student.email.ilike(like)) | (Student.student_id.ilike(like)))
    if department:
        query = query.filter(Student.department == department)
    if eligibility_status:
        query = query.filter(Student.eligibility_status == eligibility_status)
    if application_status:
        query = query.filter(Student.placement_status == application_status)
    total = query.count()
    return {"items": _paginate(query.order_by(Student.created_at.desc()), page, page_size), "total": total}


@app.get("/api/admin/students/{student_id}")
def admin_student(student_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    student = db.get(Student, student_id)
    if student is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student not found.")
    return student


@app.post("/api/admin/students")
def create_student(payload: StudentIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    data = payload.model_dump()
    email = str(data["email"]).lower().strip()
    if db.scalar(select(Student).where(Student.email == email)):
        raise HTTPException(status.HTTP_409_CONFLICT, "A student with this email already exists.")
    student = Student(**data)
    student.email = email
    db.add(student); db.commit(); db.refresh(student)
    return student


@app.patch("/api/admin/students/{student_id}")
def update_student(student_id: str, payload: StudentUpdateIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    student = db.get(Student, student_id)
    if student is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Student not found.")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(student, key, value)
    db.commit(); db.refresh(student)
    return student


@app.get("/api/admin/companies")
def admin_companies(page: int = 1, page_size: int = 10, search: str = "", authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    query = db.query(Company)
    if search:
        like = f"%{search}%"
        query = query.filter((Company.name.ilike(like)) | (Company.sector.ilike(like)))
    total = query.count()
    return {"items": _paginate(query.order_by(Company.created_at.desc()), page, page_size), "total": total}


@app.get("/api/admin/companies/{company_id}")
def admin_company(company_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    company = db.get(Company, company_id)
    if company is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Company not found.")
    return company


@app.post("/api/admin/companies")
def create_company(payload: CompanyIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    data = payload.model_dump()
    if db.scalar(select(Company).where(Company.name == data["name"])):
        raise HTTPException(status.HTTP_409_CONFLICT, "A company with this name already exists.")
    company = Company(**data)
    db.add(company); db.commit(); db.refresh(company)
    return company


@app.patch("/api/admin/companies/{company_id}")
def update_company(company_id: str, payload: CompanyUpdateIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    company = db.get(Company, company_id)
    if company is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Company not found.")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(company, key, value)
    db.commit(); db.refresh(company)
    return company


@app.get("/api/admin/eligibility")
def admin_eligibility(page: int = 1, page_size: int = 10, search: str = "", status_filter: str = "", authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    query = db.query(EligibilityRecord)
    if status_filter:
        query = query.filter(EligibilityRecord.status == status_filter)
    records = query.order_by(EligibilityRecord.updated_at.desc()).all()
    items = [_eligibility_out(record, db) for record in records]
    if search:
        lowered = search.lower()
        items = [item for item in items if lowered in " ".join(str(value) for value in item.values()).lower()]
    total = len(items)
    start = (max(page, 1) - 1) * page_size
    return {"items": items[start:start + page_size], "total": total}


@app.get("/api/admin/eligibility/records/{record_id}")
def admin_eligibility_record(record_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    record = db.get(EligibilityRecord, record_id)
    if record is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Eligibility record not found.")
    return _eligibility_out(record, db)


@app.get("/api/admin/eligibility/form-options")
def admin_eligibility_options(authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    return {"students": db.query(Student).order_by(Student.full_name).all(), "jobs": [{"id": company.id, "company_name": company.name, "title": "Open role"} for company in db.query(Company).order_by(Company.name).all()], "roles": ["Open role"], "departments": sorted({student.department for student in db.query(Student).all() if student.department}), "statuses": ["eligible", "review", "not eligible"]}


@app.post("/api/admin/eligibility")
def create_eligibility(payload: EligibilityIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    record = EligibilityRecord(**payload.model_dump())
    db.add(record); db.commit(); db.refresh(record)
    return _eligibility_out(record, db)


@app.patch("/api/admin/eligibility/{record_id}")
def update_eligibility(record_id: str, payload: EligibilityUpdateIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    record = db.get(EligibilityRecord, record_id)
    if record is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Eligibility record not found.")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(record, key, value)
    db.commit(); db.refresh(record)
    return _eligibility_out(record, db)


@app.get("/api/admin/analytics")
def admin_analytics(authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    applications = db.query(Candidate).count()
    companies = db.query(Company).count()
    companies_with_openings = db.query(Company).filter(Company.job_count > 0).count()
    shortlisted = db.query(Candidate).filter(Candidate.is_shortlisted == True).count()
    return {
        **admin_dashboard(authorization, db),
        "company_engagement_rate": round((companies_with_openings / companies) * 100) if companies else 0,
        "companies_with_openings": companies_with_openings,
        "conversion_rate": round((shortlisted / applications) * 100) if applications else 0,
        "shortlisted_applications": shortlisted,
        "student_distribution": [{"name": row.department or "Unassigned", "value": db.query(Student).filter(Student.department == row.department).count()} for row in db.query(Student.department).distinct().all()],
        "placements_by_company": [{"name": company.name, "value": company.job_count or 0} for company in db.query(Company).order_by(Company.name).limit(10).all()],
    }


@app.get("/api/admin/analytics/report")
def admin_analytics_report(authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    content = b"Job Title,Status\nSoftware Engineer,Open\n"
    headers = {"Content-Disposition": "attachment; filename=overall-report.csv"}
    return StreamingResponse(iter([content]), media_type="text/csv", headers=headers)


@app.get("/api/admin/notifications")
def admin_notifications(page: int = 1, page_size: int = 10, search: str = "", authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    query = db.query(AdminNotification)
    if search:
        like = f"%{search}%"
        query = query.filter((AdminNotification.title.ilike(like)) | (AdminNotification.message.ilike(like)))
    total = query.count()
    return {"items": _paginate(query.order_by(AdminNotification.created_at.desc()), page, page_size), "total": total}


@app.post("/api/admin/notifications")
def create_notification(payload: NotificationIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_admin_user(authorization, db)
    notification = AdminNotification(**payload.model_dump())
    db.add(notification); db.commit(); db.refresh(notification)
    return notification


def _job_out(job: Job, db: Session) -> JobOut:
    owner = db.get(User, job.owner_user_id)
    company = db.get(Company, job.company_id) if job.company_id else None
    return JobOut(
        **{key: getattr(job, key) for key in (
            "id", "owner_user_id", "company_id", "title", "department", "location",
            "employment_type", "work_mode", "experience_level", "salary_range", "openings",
            "application_deadline", "priority", "job_description", "required_skills",
            "preferred_skills", "minimum_experience_years", "education_requirement",
            "status", "shortlist_threshold", "created_at", "updated_at",
        )},
        owner_name=owner.full_name if owner else None,
        company_name=company.name if company else None,
        analysis_count=db.query(Analysis).filter(Analysis.job_id == job.id).count(),
    )


@app.get("/api/v1/jobs", response_model=list[JobOut])
def list_jobs(
    status_filter: str | None = None,
    page: int = 1,
    page_size: int = 100,
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    user = _require_recruitment_user(authorization, db)
    query = select(Job)
    if not _is_admin(user):
        query = query.where(Job.owner_user_id == user.id)
    if status_filter:
        query = query.where(Job.status == status_filter)
    jobs = db.scalars(query.order_by(Job.updated_at.desc()).offset((max(page, 1) - 1) * min(max(page_size, 1), 100)).limit(min(max(page_size, 1), 100))).all()
    return [_job_out(job, db) for job in jobs]


@app.post("/api/v1/jobs", response_model=JobOut, status_code=status.HTTP_201_CREATED)
def create_job(
    payload: JobCreateIn,
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    user = _require_recruitment_user(authorization, db)
    company = _company_for_job_request(payload.company_id, user, db)
    job = Job(
        **payload.model_dump(exclude={"company_id"}),
        owner_user_id=user.id,
        company_id=company.id if company else None,
    )
    db.add(job)
    db.flush()
    _sync_company_job_count(job.company_id, db)
    db.commit()
    db.refresh(job)
    return _job_out(job, db)


@app.get("/api/v1/jobs/{job_id}", response_model=JobOut)
def get_job(
    job_id: str,
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    user = _require_recruitment_user(authorization, db)
    return _job_out(_get_job_for_user(job_id, user, db), db)


@app.patch("/api/v1/jobs/{job_id}", response_model=JobOut)
def update_job(
    job_id: str,
    payload: JobUpdateIn,
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    user = _require_recruitment_user(authorization, db)
    job = _get_job_for_user(job_id, user, db)
    updates = payload.model_dump(exclude_unset=True)
    old_company_id = job.company_id
    if "company_id" in updates:
        company = _company_for_job_request(updates["company_id"], user, db)
        updates["company_id"] = company.id if company else None
    for key, value in updates.items():
        setattr(job, key, value)
    _sync_company_job_count(old_company_id, db)
    _sync_company_job_count(job.company_id, db)
    db.commit()
    db.refresh(job)
    return _job_out(job, db)


@app.patch("/api/v1/jobs/{job_id}/shortlist-threshold")
def update_shortlist_threshold(
    job_id: str,
    payload: ShortlistThresholdIn,
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    """Persist the score threshold used for automatic recommendations in future analyses.

    Existing candidate decisions (including manual removals) are intentionally untouched.
    """
    user = _require_recruitment_user(authorization, db)
    job = _get_job_for_user(job_id, user, db)
    job.shortlist_threshold = payload.threshold
    db.commit()
    db.refresh(job)
    return {"job_id": job.id, "threshold": job.shortlist_threshold}


@app.patch("/api/v1/jobs/{job_id}/status", response_model=JobOut)
def set_job_status(
    job_id: str,
    status_value: str = Form(..., pattern="^(draft|open|closed)$"),
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    user = _require_recruitment_user(authorization, db)
    job = _get_job_for_user(job_id, user, db)
    job.status = status_value
    _sync_company_job_count(job.company_id, db)
    db.commit()
    db.refresh(job)
    return _job_out(job, db)


@app.get("/api/v1/jobs/{job_id}/analyses", response_model=list[AnalysisOut])
def list_job_analyses(
    job_id: str,
    page: int = 1,
    page_size: int = 50,
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    user = _require_recruitment_user(authorization, db)
    job = _get_job_for_user(job_id, user, db)
    analyses = db.scalars(
        select(Analysis).where(Analysis.job_id == job.id).order_by(Analysis.created_at.desc().nullslast(), Analysis.id.desc())
        .offset((max(page, 1) - 1) * min(max(page_size, 1), 100)).limit(min(max(page_size, 1), 100))
    ).all()
    return [to_analysis_out(analysis, include_candidates=False) for analysis in analyses]


@app.post("/api/v1/analyses", response_model=AnalysisOut, status_code=status.HTTP_202_ACCEPTED)
async def create_analysis(
    background_tasks: BackgroundTasks,
    job_id: str = Form(...),
    resumes: list[UploadFile] = File(...),
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    user = _require_recruitment_user(authorization, db)
    job = _get_job_for_user(job_id, user, db)
    if job.status != "open":
        raise HTTPException(status.HTTP_409_CONFLICT, "Only open jobs can be analyzed.")
    if not resumes:
        raise HTTPException(422, "Upload at least one resume PDF.")
    if len(resumes) > settings.max_resumes_per_analysis:
        raise HTTPException(422, f"Upload at most {settings.max_resumes_per_analysis} resumes per analysis.")
    if invalid := [upload.filename or "unnamed" for upload in resumes if not (upload.filename or "").lower().endswith(".pdf")]:
        raise HTTPException(422, f"Only PDF resumes are accepted: {', '.join(invalid)}")

    analysis = Analysis(
        job_id=job.id,
        run_number=(db.scalar(select(func.max(Analysis.run_number)).where(Analysis.job_id == job.id)) or 0) + 1,
        owner_user_id=user.id,
        job_title=job.title,
        job_description=job.job_description,
        job_snapshot={
            "title": job.title,
            "department": job.department,
            "location": job.location,
            "employment_type": job.employment_type,
            "work_mode": job.work_mode,
            "experience_level": job.experience_level,
            "salary_range": job.salary_range,
            "openings": job.openings,
            "description": job.job_description,
            "required_skills": list(job.required_skills or []),
            "preferred_skills": list(job.preferred_skills or []),
            "minimum_experience_years": job.minimum_experience_years,
            "education_requirement": job.education_requirement,
            "shortlist_threshold": job.shortlist_threshold,
            "captured_at": datetime.utcnow().isoformat(),
        },
        requirements={"required": list(job.required_skills or []), "preferred": list(job.preferred_skills or [])},
        status="queued",
        total_resumes=len(resumes)
    )
    db.add(analysis)
    db.commit(); db.refresh(analysis)
    batch_dir = settings.upload_dir / analysis.id
    batch_dir.mkdir(parents=True, exist_ok=False)
    paths: list[tuple[str, str]] = []
    try:
        for upload in resumes:
            filename = Path(upload.filename or "resume.pdf").name
            destination = batch_dir / f"{uuid.uuid4()}-{filename}"
            written = 0
            signature = bytearray()
            with destination.open("wb") as output:
                while chunk := await upload.read(1024 * 1024):
                    written += len(chunk)
                    if len(signature) < 5:
                        signature.extend(chunk[: 5 - len(signature)])
                    if written > settings.max_file_size_bytes:
                        output.close(); destination.unlink(missing_ok=True)
                        raise HTTPException(413, f"{filename} exceeds the 5 MB per-file limit.")
                    output.write(chunk)
            if not bytes(signature).startswith(b"%PDF-"):
                destination.unlink(missing_ok=True)
                raise HTTPException(422, f"{filename} is not a valid PDF file.")
            paths.append((filename, str(destination)))
    except Exception:
        shutil.rmtree(batch_dir, ignore_errors=True)
        db.delete(analysis); db.commit()
        raise
    background_tasks.add_task(analyse_job, analysis.id, paths)
    return to_analysis_out(analysis, include_candidates=False)


def to_analysis_out(analysis: Analysis, include_candidates: bool = True) -> AnalysisOut:
    candidates = sorted(analysis.candidates, key=lambda item: item.overall_score, reverse=True) if include_candidates else []
    # Handle both old list format and new dict format for backward compatibility
    requirements = analysis.requirements if isinstance(analysis.requirements, dict) else {"all": analysis.requirements or [], "required": [], "preferred": []}
    created_at = analysis.created_at
    if created_at is not None and created_at.tzinfo is None:
        # Historical rows were stored as naive UTC via datetime.utcnow().
        created_at = created_at.replace(tzinfo=timezone.utc)
    return AnalysisOut(id=analysis.id, job_id=analysis.job_id, run_number=analysis.run_number,
                       job_title=analysis.job_title, requirements=requirements, status=analysis.status,
                       error=analysis.error, created_at=created_at, completed_at=analysis.completed_at,
                       job_snapshot=analysis.job_snapshot,
                       candidate_count=len(analysis.candidates), candidates=[CandidateOut.model_validate(item) for item in candidates])


def _to_export_string(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(item) for item in value if item)
    if isinstance(value, dict):
        return ", ".join(f"{key}: {val}" for key, val in value.items())
    return str(value)


def _excel_safe(value):
    """Prevent user supplied text from being interpreted as an Excel formula."""
    if isinstance(value, str) and (value.startswith(("\t", "\r", "\n")) or value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@"))):
        return "'" + value
    return value


def _export_candidate_fields(candidate) -> dict:
    """Recover resume metadata from stored JSON or raw text when earlier stages dropped it."""
    entities = getattr(candidate, "entities", {}) or {}
    fallback_entities = extract_entities(getattr(candidate, "text", "") or "")

    email = candidate.email or _to_export_string(entities.get("emails") or fallback_entities.get("emails") or [])
    phone = _to_export_string(entities.get("phones") or fallback_entities.get("phones") or [])
    location = _to_export_string(entities.get("locations") or fallback_entities.get("locations") or [])
    education = candidate.education or extract_education(getattr(candidate, "text", "") or "") or "Not specified"

    return {
        "email": email,
        "phone": phone,
        "location": location,
        "education": education,
    }


def _add_resume_link(cell, candidate: Candidate) -> None:
    """Link to the frontend's protected resume viewer, which uses the user's session."""
    cell.value = "View resume"
    frontend_url = (settings.frontend_base_url or next(iter(settings.allowed_origins), "http://localhost:5173")).rstrip("/")
    cell.hyperlink = f"{frontend_url}/dashboard/resume/{candidate.id}"
    cell.style = "Hyperlink"


def _candidate_resume_path(candidate: Candidate) -> Path | None:
    batch_dir = settings.upload_dir / candidate.analysis_id
    if not batch_dir.exists():
        return None
    if candidate.stored_filename:
        exact_file = batch_dir / Path(candidate.stored_filename).name
        if exact_file.is_file():
            return exact_file
    legacy_matches = [
        file_path for file_path in batch_dir.glob("*.pdf")
        if file_path.name.endswith(f"-{candidate.filename}") or file_path.name == candidate.filename
    ]
    return legacy_matches[0] if len(legacy_matches) == 1 else None


def build_report_workbook(analysis: Analysis) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Job title", _excel_safe(analysis.job_title or "Untitled role")])
    summary.append(["Job ID", analysis.job_id or ""])
    summary.append(["Candidates", len(analysis.candidates)])

    if isinstance(analysis.requirements, dict):
        required = analysis.requirements.get("required", [])
        preferred = analysis.requirements.get("preferred", [])
        req_text = f"Required: {', '.join(required)} | Preferred: {', '.join(preferred)}"
    else:
        req_text = ", ".join(analysis.requirements or [])
    summary.append(["Requirements", _excel_safe(req_text)])
    summary.append([])
    summary.append(["Candidate ID", "Candidate", "Email", "Phone", "Location", "Education", "Top Skills", "Shortlisted", "Resume"])
    for candidate in sorted(analysis.candidates, key=lambda item: item.overall_score, reverse=True):
        contact_fields = _export_candidate_fields(candidate)
        summary.append([
            candidate.id, _excel_safe(candidate.name), _excel_safe(contact_fields["email"]), _excel_safe(contact_fields["phone"]),
            _excel_safe(contact_fields["location"]), _excel_safe(contact_fields["education"]),
            _excel_safe(", ".join(candidate.skills[:8])), bool(candidate.is_shortlisted),
            None,
        ])
        _add_resume_link(summary.cell(summary.max_row, 9), candidate)

    ranking = workbook.create_sheet("Candidate ranking")
    ranking.append([
        "Rank", "Candidate ID", "Candidate", "Email", "Overall", "Semantic", "Keyword", "Experience",
        "Recommendation", "Skills", "Missing skills", "Insight", "Shortlisted", "Resume"
    ])

    for rank, candidate in enumerate(sorted(analysis.candidates, key=lambda item: item.overall_score, reverse=True), 1):
        contact_fields = _export_candidate_fields(candidate)
        ranking.append([
            rank,
            candidate.id, _excel_safe(candidate.name), _excel_safe(contact_fields["email"]),
            candidate.overall_score,
            candidate.semantic_score,
            candidate.keyword_score,
            candidate.experience_score,
            _excel_safe(candidate.recommendation), _excel_safe(", ".join(candidate.skills)),
            _excel_safe(", ".join(candidate.missing_skills)), _excel_safe(candidate.insight),
            bool(candidate.is_shortlisted),
            None,
        ])
        _add_resume_link(ranking.cell(ranking.max_row, 14), candidate)

    details = workbook.create_sheet("Candidate details")
    details.append([
        "Rank", "Candidate ID", "Candidate", "Email", "Phone", "Overall", "Semantic", "Keyword", "Experience",
        "Recommendation", "Experience years", "Education", "Skills", "Missing skills",
        "Projects", "Experience details", "Organizations", "Locations", "Insight", "Shortlisted", "Resume"
    ])
    for rank, candidate in enumerate(sorted(analysis.candidates, key=lambda item: item.overall_score, reverse=True), 1):
        entities = getattr(candidate, "entities", {}) or {}
        contact_fields = _export_candidate_fields(candidate)
        details.append([
            rank,
            candidate.id, _excel_safe(candidate.name), _excel_safe(contact_fields["email"]), _excel_safe(contact_fields["phone"]),
            candidate.overall_score,
            candidate.semantic_score,
            candidate.keyword_score,
            candidate.experience_score,
            _excel_safe(candidate.recommendation),
            entities.get("experience_years") if isinstance(entities.get("experience_years"), (int, float)) else 0,
            contact_fields["education"],
            _excel_safe(", ".join(candidate.skills)),
            _excel_safe(", ".join(candidate.missing_skills)),
            _excel_safe(_to_export_string(getattr(candidate, "projects", []) or [])),
            _excel_safe(_to_export_string(getattr(candidate, "experience_details", []) or [])),
            _excel_safe(_to_export_string(entities.get("organizations") or [])),
            _excel_safe(contact_fields["location"]),
            _excel_safe(getattr(candidate, "insight", "")), bool(candidate.is_shortlisted),
            None,
        ])
        _add_resume_link(details.cell(details.max_row, 21), candidate)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
            sheet.column_dimensions[column[0].column_letter].width = width

    return workbook


@app.get("/api/v1/analyses", response_model=list[AnalysisOut])
def list_analyses(page: int = 1, page_size: int = 50, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    user = _require_recruitment_user(authorization, db)
    query = select(Analysis).order_by(Analysis.created_at.desc().nullslast(), Analysis.id.desc())
    if not _is_admin(user): query = query.where(Analysis.owner_user_id == user.id)
    analyses = db.scalars(query.offset((max(page, 1) - 1) * min(max(page_size, 1), 100)).limit(min(max(page_size, 1), 100))).all()
    return [to_analysis_out(analysis, include_candidates=False) for analysis in analyses]


@app.get("/api/v1/analyses/{analysis_id}", response_model=AnalysisOut)
def get_analysis(analysis_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    user = _require_recruitment_user(authorization, db)
    analysis = _get_analysis_for_user(analysis_id, user, db, include_candidates=True)
    return to_analysis_out(analysis)


@app.delete("/api/v1/analyses/{analysis_id}")
def delete_analysis(analysis_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    """Permanently delete one owned analysis and its analysis-specific resume files."""
    user = _require_recruitment_user(authorization, db)
    analysis = _get_analysis_for_user(analysis_id, user, db, include_candidates=True)
    if analysis.status in {"queued", "processing"}:
        raise HTTPException(status.HTTP_409_CONFLICT, "This analysis is still running. Wait for it to finish before deleting it.")

    upload_root = settings.upload_dir.resolve()
    batch_dir = settings.upload_dir / analysis.id
    staged_dir = upload_root / f".deleting-{analysis.id}-{uuid.uuid4().hex}"
    staged = False
    try:
        resolved_batch = batch_dir.resolve()
        if resolved_batch.parent != upload_root:
            raise HTTPException(status.HTTP_409_CONFLICT, "Analysis files are outside the managed upload directory and were not deleted.")
        if batch_dir.exists():
            batch_dir.rename(staged_dir)
            staged = True

        db.delete(analysis)
        db.commit()
    except HTTPException:
        db.rollback()
        if staged and staged_dir.exists() and not batch_dir.exists():
            staged_dir.rename(batch_dir)
        raise
    except Exception as exc:
        db.rollback()
        if staged and staged_dir.exists() and not batch_dir.exists():
            try:
                staged_dir.rename(batch_dir)
            except OSError:
                logger.exception("Could not restore resume directory for analysis %s after database rollback", analysis_id)
        logger.exception("Could not delete analysis %s", analysis_id)
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, "Analysis deletion failed. No analysis records were removed.") from exc

    files_deleted = True
    if staged:
        try:
            shutil.rmtree(staged_dir)
        except OSError:
            files_deleted = False
            logger.exception("Analysis %s was deleted, but its staged resume files could not be removed", analysis_id)
    return {"id": analysis_id, "deleted": True, "files_deleted": files_deleted}


@app.get("/api/v1/analyses/{analysis_id}/status", response_model=AnalysisStatusOut)
def get_analysis_status(analysis_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    user = _require_recruitment_user(authorization, db)
    analysis = _get_analysis_for_user(analysis_id, user, db, include_candidates=True)
    
    completed = len(analysis.candidates)
    
    completed_resumes = [
        CompletedResumeOut(
            id=c.id,
            filename=c.filename,
            candidate_name=c.name,
            ats_score=c.overall_score
        )
        for c in analysis.candidates
    ]
    
    current_resume = None
    if analysis.current_resume_filename:
        current_resume = CurrentResumeState(
            filename=analysis.current_resume_filename,
            candidate_name=analysis.current_resume_candidate_name
        )
        
    return AnalysisStatusOut(
        analysis_id=analysis.id,
        status=analysis.status,
        phase=analysis.current_phase,
        current_stage=analysis.current_stage,
        total=analysis.total_resumes,
        completed=completed,
        phase_completed=analysis.phase_completed,
        phase_total=analysis.phase_total,
        current_resume=current_resume,
        completed_resumes=completed_resumes,
        started_at=analysis.created_at,
        current_resume_started_at=analysis.current_resume_started_at,
        completed_at=analysis.completed_at
        ,error=analysis.error
    )


@app.get("/api/v1/analyses/{analysis_id}/report.xlsx")
def download_report(analysis_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    user = _require_recruitment_user(authorization, db)
    analysis = _get_analysis_for_user(analysis_id, user, db, include_candidates=True)
    if analysis.status != "completed": raise HTTPException(409, "Analysis is not complete.")

    workbook = build_report_workbook(analysis)
    from io import BytesIO
    buffer = BytesIO(); workbook.save(buffer); buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="hiresense-ai-{analysis_id}.xlsx"'})


@app.get("/api/v1/candidates/{candidate_id}/resume")
def download_candidate_resume(candidate_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    """Download the original resume PDF for a specific candidate."""
    user = _require_recruitment_user(authorization, db)
    candidate = db.scalar(select(Candidate).where(Candidate.id == candidate_id))
    if not candidate: raise HTTPException(404, "Candidate not found.")
    _get_analysis_for_user(candidate.analysis_id, user, db)
    
    matching_file = _candidate_resume_path(candidate)
    
    if not matching_file: raise HTTPException(404, "Resume file not found.")
    
    return FileResponse(
        path=str(matching_file),
        media_type="application/pdf",
        filename=candidate.filename,
        headers={"Content-Disposition": f'inline; filename="{candidate.filename}"'}
    )


@app.post("/api/v1/job-descriptions/extract", response_model=JobDescriptionExtractOut)
async def extract_job_description(file: UploadFile = File(...), authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    _require_recruitment_user(authorization, db)
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(422, "Only PDF files are supported.")
    
    # Save to a temporary file
    temp_dir = settings.upload_dir / "temp_jd"
    temp_dir.mkdir(parents=True, exist_ok=True)
    temp_path = temp_dir / f"{uuid.uuid4()}-{file.filename}"
    
    try:
        written = 0
        signature = bytearray()
        with temp_path.open("wb") as output:
            while chunk := await file.read(1024 * 1024):
                written += len(chunk)
                if len(signature) < 5:
                    signature.extend(chunk[: 5 - len(signature)])
                if written > settings.max_file_size_bytes:
                    raise HTTPException(413, "File exceeds the size limit.")
                output.write(chunk)
        if not bytes(signature).startswith(b"%PDF-"):
            raise HTTPException(422, "Unable to read the PDF.")
        
        # We need pages to return in schema.
        import fitz
        pages = 0
        with fitz.open(temp_path) as doc:
            pages = len(doc)

        from app.nlp import extract_jd_pdf_text
        text = extract_jd_pdf_text(str(temp_path))
        
        # Validation
        if not text or len(text) < 50 or not any(c.isalpha() for c in text):
            raise HTTPException(422, "No readable job description text was found in this PDF.")
            
        return JobDescriptionExtractOut(text=text, filename=file.filename or "jd.pdf", pages=pages)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(422, "Unable to read the PDF.")
    finally:
        temp_path.unlink(missing_ok=True)


# ============================================================================
# SHORTLISTING AND EMAIL ENDPOINTS
# ============================================================================

@app.post("/api/v1/candidates/{candidate_id}/shortlist", response_model=CandidateOut)
def update_shortlist(candidate_id: str, update: ShortlistUpdateIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    """Update shortlist status for a candidate."""
    user = _require_recruitment_user(authorization, db)
    candidate = db.scalar(select(Candidate).where(Candidate.id == candidate_id))
    if not candidate:
        raise HTTPException(404, "Candidate not found.")
    _get_analysis_for_user(candidate.analysis_id, user, db)
    
    candidate.is_shortlisted = update.is_shortlisted
    candidate.shortlist_source = "manual"
    if update.notes is not None:
        candidate.notes = update.notes
    
    db.commit()
    db.refresh(candidate)
    return CandidateOut.model_validate(candidate)


@app.post("/api/v1/analyses/{analysis_id}/shortlist/bulk-remove")
def bulk_remove_from_shortlist(
    analysis_id: str,
    payload: CandidateIdsIn,
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    user = _require_recruitment_user(authorization, db)
    analysis = _get_analysis_for_user(analysis_id, user, db)
    candidate_ids = set(payload.candidate_ids)
    candidates = db.scalars(select(Candidate).where(
        Candidate.analysis_id == analysis.id, Candidate.id.in_(candidate_ids)
    )).all()
    if len(candidates) != len(candidate_ids):
        raise HTTPException(404, "One or more selected candidates do not belong to this analysis.")
    changed = 0
    for candidate in candidates:
        if candidate.is_shortlisted:
            candidate.is_shortlisted = False
            candidate.shortlist_source = "manual"
            changed += 1
    db.commit()
    return {"updated_count": changed, "candidate_ids": list(candidate_ids)}


@app.post("/api/v1/analyses/{analysis_id}/shortlist/bulk-add")
def bulk_add_to_shortlist(
    analysis_id: str,
    payload: CandidateIdsIn,
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    user = _require_recruitment_user(authorization, db)
    analysis = _get_analysis_for_user(analysis_id, user, db)
    candidate_ids = set(payload.candidate_ids)
    candidates = db.scalars(select(Candidate).where(
        Candidate.analysis_id == analysis.id, Candidate.id.in_(candidate_ids)
    )).all()
    if len(candidates) != len(candidate_ids):
        raise HTTPException(404, "One or more selected candidates do not belong to this analysis.")
    changed = 0
    for candidate in candidates:
        if not candidate.is_shortlisted or candidate.shortlist_source != "manual":
            candidate.is_shortlisted = True
            candidate.shortlist_source = "manual"
            changed += 1
    db.commit()
    return {"updated_count": changed, "candidate_ids": list(candidate_ids)}


@app.post("/api/v1/analyses/{analysis_id}/candidates/export.xlsx")
def export_analysis_candidates(
    analysis_id: str,
    payload: CandidateIdsIn,
    authorization: str | None = Header(None, alias="Authorization"),
    db: Session = Depends(get_db),
):
    """Export exactly the validated candidates supplied by the shortlist UI."""
    user = _require_recruitment_user(authorization, db)
    analysis = _get_analysis_for_user(analysis_id, user, db)
    candidate_ids = set(payload.candidate_ids)
    candidates = db.scalars(select(Candidate).where(
        Candidate.analysis_id == analysis.id, Candidate.id.in_(candidate_ids)
    ).order_by(Candidate.overall_score.desc(), Candidate.name.asc())).all()
    if len(candidates) != len(candidate_ids):
        raise HTTPException(404, "One or more selected candidates do not belong to this analysis.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Candidates"
    sheet.append(["Candidate ID", "Candidate", "Email", "Phone", "Match Score", "Skills", "Experience (years)", "Education", "Shortlist status", "Shortlist source", "Job", "Analysis ID", "Resume"])
    for candidate in candidates:
        contact = _export_candidate_fields(candidate)
        entities = candidate.entities or {}
        sheet.append([
            candidate.id, _excel_safe(candidate.name), _excel_safe(contact["email"]), _excel_safe(contact["phone"]),
            candidate.overall_score, _excel_safe(", ".join(candidate.skills or [])),
            entities.get("experience_years", ""), _excel_safe(contact["education"]),
            "Shortlisted" if candidate.is_shortlisted else "Not shortlisted", candidate.shortlist_source,
            _excel_safe(analysis.job_title or ""), analysis.id, None,
        ])
        resume_cell = sheet.cell(sheet.max_row, 13)
        if _candidate_resume_path(candidate):
            _add_resume_link(resume_cell, candidate)
        else:
            resume_cell.value = "Unavailable"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        sheet.column_dimensions[column[0].column_letter].width = width
    from io import BytesIO
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={
        "Content-Disposition": f'attachment; filename="hiresense-ai-candidates-{analysis.id}.xlsx"'
    })


@app.get("/api/v1/analyses/{analysis_id}/shortlisted", response_model=list[ShortlistedCandidateOut])
def get_shortlisted_candidates(analysis_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    """Get all shortlisted candidates for an analysis."""
    user = _require_recruitment_user(authorization, db)
    _get_analysis_for_user(analysis_id, user, db)
    candidates = db.scalars(
        select(Candidate)
        .where(Candidate.analysis_id == analysis_id)
        .where(Candidate.is_shortlisted == True)
        .order_by(Candidate.overall_score.desc())
    ).all()
    
    return [ShortlistedCandidateOut.model_validate(c) for c in candidates]


@app.post("/api/v1/analyses/{analysis_id}/filtered-candidates", response_model=list[CandidateOut])
def get_filtered_candidates(analysis_id: str, filters: FilterCriteriaIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    """Get candidates with applied filters."""
    user = _require_recruitment_user(authorization, db)
    _get_analysis_for_user(analysis_id, user, db)
    query = select(Candidate).where(Candidate.analysis_id == analysis_id)
    
    if filters.min_score is not None:
        query = query.where(Candidate.overall_score >= filters.min_score)
    
    if filters.max_score is not None:
        query = query.where(Candidate.overall_score <= filters.max_score)
    
    if filters.recommendation:
        query = query.where(Candidate.recommendation == filters.recommendation)
    
    if filters.is_shortlisted is not None:
        query = query.where(Candidate.is_shortlisted == filters.is_shortlisted)
    
    if filters.email_sent is not None:
        query = query.where(Candidate.email_sent == filters.email_sent)
    
    query = query.order_by(Candidate.overall_score.desc())
    candidates = db.scalars(query).all()
    
    # Filter by skills if provided
    if filters.skills_filter:
        filtered_list = []
        for candidate in candidates:
            candidate_skills = set(s.lower() for s in (candidate.skills or []))
            search_skills = set(s.lower() for s in filters.skills_filter)
            if search_skills.issubset(candidate_skills):
                filtered_list.append(candidate)
        candidates = filtered_list
    
    return [CandidateOut.model_validate(c) for c in candidates]


@app.post("/api/v1/candidates/send-email", response_model=EmailResponseOut)
def send_email_to_candidates(email_request: SendEmailIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    """Send email to selected candidates."""
    user = _require_recruitment_user(authorization, db)
    
    # Get candidates
    candidate_query = select(Candidate).join(Analysis, Candidate.analysis_id == Analysis.id)
    if not _is_admin(user):
        candidate_query = candidate_query.where(Analysis.owner_user_id == user.id)
    if email_request.recipient_ids is not None:
        if not email_request.recipient_ids:
            raise HTTPException(422, "Select at least one recipient.")
        candidates = db.scalars(candidate_query.where(Candidate.id.in_(email_request.recipient_ids))).all()
        if len(candidates) != len(set(email_request.recipient_ids)):
            raise HTTPException(404, "One or more selected candidates were not found or are not accessible.")
    else:
        # Send to all shortlisted
        candidates = db.scalars(candidate_query.where(Candidate.is_shortlisted == True)).all()
    
    if not candidates:
        return EmailResponseOut(
            success=False,
            message="No candidates found to send emails to.",
            sent_count=0,
            failed_count=0
        )
    for candidate in candidates:
        _get_analysis_for_user(candidate.analysis_id, user, db)
    
    email_payloads = []
    seen_emails = set()
    errors = []
    failed_count = 0
    skipped_count = 0
    for candidate in candidates:
        if not candidate.email:
            skipped_count += 1
            errors.append({"candidate_id": candidate.id, "candidate_name": candidate.name, "error": "No email address on file", "status": "skipped"})
            continue

        normalized_email = candidate.email.strip().lower()
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", normalized_email):
            skipped_count += 1
            errors.append({"candidate_id": candidate.id, "candidate_name": candidate.name, "error": "Invalid email address", "status": "skipped"})
            continue
        if normalized_email in seen_emails:
            skipped_count += 1
            errors.append({"candidate_id": candidate.id, "candidate_name": candidate.name, "error": "Duplicate recipient email skipped", "status": "skipped"})
            continue
        seen_emails.add(normalized_email)

        if email_request.template_type == "procedure":
            template = EmailService.get_procedure_email_template(candidate.name)
        elif email_request.template_type == "rejection":
            template = EmailService.get_rejection_email_template(candidate.name)
        elif email_request.template_type == "custom":
            template = {
                "subject": email_request.custom_subject or "Message from HireSense AI",
                "body": email_request.custom_body or "",
                "html_body": email_request.custom_html_body
            }
        else:
            continue

        email_payloads.append({
            "candidate_id": candidate.id,
            "candidate_name": candidate.name,
            "recipient_email": candidate.email,
            "subject": template["subject"],
            "body": template["body"],
            "html_body": template.get("html_body"),
        })

    result = EmailService.send_emails(email_payloads)
    sent_count = 0
    result_errors = result.get("errors", [])
    if email_payloads and not result.get("success") and not result_errors:
        message = result.get("message", "Email delivery failed or SMTP is not configured.")
        errors.extend({"candidate_id": item["candidate_id"], "candidate_name": item["candidate_name"], "error": message} for item in email_payloads)
        result_errors = [{"candidate_email": item["recipient_email"], "error": message} for item in email_payloads]

    for item in email_payloads:
        candidate = db.scalar(select(Candidate).where(Candidate.id == item["candidate_id"]))
        if not candidate:
            continue

        if item["recipient_email"] in [error.get("candidate_email") for error in result_errors]:
            failed_count += 1
            if not any(error.get("candidate_id") == candidate.id for error in errors):
                errors.append({"candidate_id": candidate.id, "candidate_name": candidate.name, "error": next((error["error"] for error in result_errors if error.get("candidate_email") == item["recipient_email"]), "Email delivery failed")})
            continue

        candidate.email_sent = True
        candidate.email_sent_at = datetime.utcnow()
        sent_count += 1

    if email_request.recipient_ids is None:
        for candidate in candidates:
            if candidate.email and candidate.email_sent:
                candidate.email_sent = True

    db.commit()

    if not email_payloads:
        return EmailResponseOut(
            success=False,
            message="No valid recipient email addresses were available to send.",
            sent_count=0,
            failed_count=failed_count,
            skipped_count=skipped_count,
            errors=errors
        )

    return EmailResponseOut(
        success=sent_count > 0,
        message=(result.get("message") if sent_count == 0 and result.get("message") else f"Sent {sent_count} emails successfully" + (f"; {failed_count} failed" if failed_count > 0 else "")),
        sent_count=sent_count,
        failed_count=failed_count,
        skipped_count=skipped_count,
        errors=errors
    )


@app.post("/api/v1/candidates/{candidate_id}/schedule-call", response_model=CandidateOut)
def schedule_call(candidate_id: str, call_request: CallProcedureIn, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    """Schedule a call/procedure for a candidate."""
    user = _require_recruitment_user(authorization, db)
    candidate = db.scalar(select(Candidate).where(Candidate.id == candidate_id))
    if not candidate:
        raise HTTPException(404, "Candidate not found.")
    _get_analysis_for_user(candidate.analysis_id, user, db)
    
    candidate.call_scheduled = True
    candidate.call_scheduled_at = call_request.scheduled_at or datetime.utcnow()
    if call_request.notes:
        candidate.notes = call_request.notes
    
    db.commit()
    db.refresh(candidate)
    return CandidateOut.model_validate(candidate)


@app.get("/api/v1/analyses/{analysis_id}/shortlisted-report.xlsx")
def download_shortlisted_report(analysis_id: str, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    """Download Excel report of only shortlisted candidates."""
    user = _require_recruitment_user(authorization, db)
    analysis = _get_analysis_for_user(analysis_id, user, db, include_candidates=True)
    
    # Filter only shortlisted candidates
    shortlisted = [c for c in analysis.candidates if c.is_shortlisted]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Shortlisted Candidates"
    
    # Headers
    sheet.append([
        "Rank", "Candidate ID", "Candidate", "Email", "Phone", "Overall Score", "Semantic", "Keyword", "Experience",
        "Recommendation", "Education", "Skills", "Missing Skills", "Experience Years", "Location", "Resume"
    ])
    
    # Data rows
    for rank, candidate in enumerate(sorted(shortlisted, key=lambda c: c.overall_score, reverse=True), 1):
        contact_fields = _export_candidate_fields(candidate)
        entities = getattr(candidate, "entities", {}) or {}
        
        sheet.append([
            rank,
            candidate.id, _excel_safe(candidate.name), _excel_safe(contact_fields["email"]), _excel_safe(contact_fields["phone"]),
            candidate.overall_score,
            candidate.semantic_score,
            candidate.keyword_score,
            candidate.experience_score,
            _excel_safe(candidate.recommendation), _excel_safe(contact_fields["education"]),
            _excel_safe(", ".join(candidate.skills)), _excel_safe(", ".join(candidate.missing_skills)),
            entities.get("experience_years", ""),
            _excel_safe(contact_fields["location"]),
            None,
        ])
        _add_resume_link(sheet.cell(sheet.max_row, 16), candidate)
    
    # Format sheet
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        sheet.column_dimensions[column[0].column_letter].width = width
    
    from io import BytesIO
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="hiresense-ai-shortlisted-{analysis_id}.xlsx"'}
    )


@app.get("/api/v1/candidates/strong-matches")
def get_strong_matches(analysis_id: str = None, authorization: str | None = Header(None, alias="Authorization"), db: Session = Depends(get_db)):
    """Get all 'Strong Match' candidates for outreach."""
    user = _require_recruitment_user(authorization, db)
    query = select(Candidate).join(Analysis, Candidate.analysis_id == Analysis.id).where(Candidate.recommendation == "Strong Match")
    if not _is_admin(user): query = query.where(Analysis.owner_user_id == user.id)
    if analysis_id:
        _get_analysis_for_user(analysis_id, user, db)
        query = query.where(Candidate.analysis_id == analysis_id)
    candidates = db.scalars(query.order_by(Candidate.overall_score.desc())).all()
    
    return [CandidateOut.model_validate(c) for c in candidates]
