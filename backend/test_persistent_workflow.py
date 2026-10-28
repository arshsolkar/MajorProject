"""Isolated API and migration checks for the persistent recruitment workflow."""
import asyncio
import os
import tempfile
import unittest
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path

_temp = tempfile.TemporaryDirectory(prefix="hiresense-workflow-")
os.environ["DATABASE_URL"] = f"sqlite:///{Path(_temp.name) / 'workflow.db'}"
os.environ["UPLOAD_DIR"] = str(Path(_temp.name) / "uploads")

from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, inspect, text

from app import main
from app.config import settings
from app.database import Base, SessionLocal, engine
from app.models import Analysis, Candidate, Job, User
from sqlalchemy.orm import Session
from app.services import automatically_shortlisted


class PersistentWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        Base.metadata.create_all(engine)
        cls.db = SessionLocal()
        cls.recruiter = User(full_name="Workflow Recruiter", email="workflow@example.test", password_hash="x", module="recruiter", role="recruiter", is_approved=True, company_name="Workflow Co")
        cls.other = User(full_name="Other Recruiter", email="other@example.test", password_hash="x", module="recruiter", role="recruiter", is_approved=True, company_name="Other Co")
        cls.admin = User(full_name="Workflow Admin", email="admin-workflow@example.test", password_hash="x", module="admin_tpo", role="tpo", is_approved=True)
        cls.db.add_all([cls.recruiter, cls.other, cls.admin]); cls.db.commit()
        cls.tokens = {u.id: main._make_token(u).access_token for u in [cls.recruiter, cls.other, cls.admin]}
        main.analyse_job = lambda *_args: None

    @classmethod
    def tearDownClass(cls):
        cls.db.close()
        _temp.cleanup()

    def request(self, method, path, token=None, **kwargs):
        async def run():
            headers = {"Authorization": f"Bearer {token}"} if token else {}
            async with AsyncClient(transport=ASGITransport(app=main.app), base_url="http://test") as client:
                return await client.request(method, path, headers=headers, **kwargs)
        return asyncio.run(run())

    def test_shortlist_threshold_uses_normalized_zero_to_hundred_scores(self):
        self.assertTrue(automatically_shortlisted(70, 70))
        self.assertTrue(automatically_shortlisted(91, 70))
        self.assertFalse(automatically_shortlisted(69, 70))

    def test_analysis_history_sorts_by_creation_time_and_delete_is_permanent_and_authorized(self):
        owner_token = self.tokens[self.recruiter.id]
        job_response = self.request("POST", "/api/v1/jobs", owner_token, json={
            "title": "Deletion Integrity Role", "job_description": "A role used for analysis deletion tests.",
            "required_skills": ["Python"],
        })
        self.assertEqual(job_response.status_code, 201, job_response.text)
        job_id = job_response.json()["id"]
        created = datetime(2026, 10, 28, 15, 0, 0)
        older = Analysis(id="zzzz-older-analysis", job_id=job_id, owner_user_id=self.recruiter.id,
                         job_title="Deletion Integrity Role", job_description="old snapshot", requirements={},
                         status="completed", run_number=2, created_at=created, completed_at=datetime(2026, 10, 28, 17, 0, 0))
        newer = Analysis(id="aaaa-newer-analysis", job_id=job_id, owner_user_id=self.recruiter.id,
                         job_title="Deletion Integrity Role", job_description="new snapshot", requirements={},
                         status="completed", run_number=3, created_at=created + timedelta(minutes=20), completed_at=created + timedelta(minutes=21))
        active = Analysis(id="active-analysis", job_id=job_id, owner_user_id=self.recruiter.id,
                          job_title="Deletion Integrity Role", job_description="active snapshot", requirements={},
                          status="processing", run_number=1, created_at=created - timedelta(minutes=10))
        older_candidate = Candidate(analysis_id=older.id, filename="older.pdf", stored_filename="older.pdf",
                                    name="Older Candidate", email="older@example.test", text="old resume",
                                    skills=["Python"], missing_skills=[], entities={}, projects=[], education="BS",
                                    experience_details=[], semantic_score=70, keyword_score=70, experience_score=70,
                                    overall_score=70, recommendation="Good Match", insight="", is_shortlisted=True,
                                    shortlist_source="manual")
        newer_candidate = Candidate(analysis_id=newer.id, filename="newer.pdf", stored_filename="newer.pdf",
                                    name="Newer Candidate", email="newer@example.test", text="new resume",
                                    skills=["Python"], missing_skills=[], entities={}, projects=[], education="BS",
                                    experience_details=[], semantic_score=90, keyword_score=90, experience_score=90,
                                    overall_score=90, recommendation="Strong Match", insight="", is_shortlisted=True,
                                    shortlist_source="automatic")
        self.db.add_all([older, newer, active, older_candidate, newer_candidate]); self.db.commit()
        older_id, newer_id, active_id = older.id, newer.id, active.id
        older_candidate_id, newer_candidate_id = older_candidate.id, newer_candidate.id
        old_dir = settings.upload_dir / older_id
        new_dir = settings.upload_dir / newer_id
        old_dir.mkdir(parents=True, exist_ok=True); new_dir.mkdir(parents=True, exist_ok=True)
        (old_dir / "older.pdf").write_bytes(b"old private resume")
        (new_dir / "newer.pdf").write_bytes(b"new private resume")
        cache_dir = settings.upload_dir / ".extraction_cache"
        cache_dir.mkdir(exist_ok=True)
        shared_cache = cache_dir / "shared-test-cache.txt"
        shared_cache.write_text("shared cache", encoding="utf-8")

        history = self.request("GET", "/api/v1/analyses", owner_token).json()
        ordered = [item for item in history if item["job_id"] == job_id]
        self.assertEqual([item["id"] for item in ordered], [newer_id, older_id, active_id])
        self.assertEqual([item["run_number"] for item in ordered], [3, 2, 1])
        self.assertEqual(datetime.fromisoformat(ordered[0]["created_at"]).replace(tzinfo=None), newer.created_at)
        self.assertLess(datetime.fromisoformat(ordered[0]["completed_at"]).replace(tzinfo=None), datetime.fromisoformat(ordered[1]["completed_at"]).replace(tzinfo=None))
        selected = self.request("GET", f"/api/v1/analyses/{newer_id}", owner_token)
        self.assertEqual(selected.json()["candidates"][0]["id"], newer_candidate_id)
        self.assertEqual(self.request("DELETE", f"/api/v1/analyses/{active_id}", owner_token).status_code, 409)
        self.assertTrue((settings.upload_dir / active_id).exists() is False)

        self.assertEqual(self.request("DELETE", f"/api/v1/analyses/{older_id}", self.tokens[self.other.id]).status_code, 404)
        self.assertEqual(self.request("DELETE", f"/api/v1/analyses/{older_id}").status_code, 401)

        # A failed database commit must restore the staged files and keep all rows.
        original_commit = Session.commit
        failed_once = {"value": False}
        def fail_delete_commit(session):
            if not failed_once["value"]:
                failed_once["value"] = True
                raise RuntimeError("simulated commit failure")
            return original_commit(session)
        Session.commit = fail_delete_commit
        try:
            failed_delete = self.request("DELETE", f"/api/v1/analyses/{older_id}", owner_token)
        finally:
            Session.commit = original_commit
        self.assertEqual(failed_delete.status_code, 500)
        self.db.expire_all()
        self.assertIsNotNone(self.db.get(Analysis, older_id))
        self.assertIsNotNone(self.db.get(Candidate, older_candidate_id))
        self.assertTrue((old_dir / "older.pdf").is_file())

        admin_delete = self.request("DELETE", f"/api/v1/analyses/{older_id}", self.tokens[self.admin.id])
        self.assertEqual(admin_delete.status_code, 200, admin_delete.text)
        self.assertTrue(admin_delete.json()["deleted"])
        self.assertTrue(admin_delete.json()["files_deleted"])
        self.db.expire_all()
        self.assertIsNone(self.db.get(Analysis, older_id))
        self.assertIsNone(self.db.get(Candidate, older_candidate_id))
        self.assertFalse(old_dir.exists())
        self.assertTrue(new_dir.exists())
        self.assertTrue(shared_cache.is_file())
        self.assertEqual(self.request("GET", f"/api/v1/analyses/{older_id}", owner_token).status_code, 404)
        self.assertEqual(self.request("GET", f"/api/v1/analyses/{older_id}/report.xlsx", owner_token).status_code, 404)
        self.assertEqual(self.request("GET", f"/api/v1/candidates/{older_candidate_id}/resume", owner_token).status_code, 404)
        remaining = self.request("GET", f"/api/v1/jobs/{job_id}/analyses", self.tokens[self.admin.id]).json()
        self.assertEqual([item["id"] for item in remaining], [newer_id, active_id])
        self.assertEqual(self.request("GET", f"/api/v1/jobs/{job_id}", owner_token).json()["analysis_count"], 2)

        final_delete = self.request("DELETE", f"/api/v1/analyses/{newer_id}", owner_token)
        self.assertEqual(final_delete.status_code, 200, final_delete.text)
        self.db.expire_all()
        self.assertIsNone(self.db.get(Analysis, newer_id))
        self.assertIsNone(self.db.get(Candidate, newer_candidate_id))
        self.assertEqual(self.request("GET", f"/api/v1/analyses/{newer_id}", self.tokens[self.admin.id]).status_code, 404)
        self.assertEqual(self.request("GET", f"/api/v1/jobs/{job_id}", owner_token).json()["analysis_count"], 1)
        self.assertEqual(self.request("DELETE", f"/api/v1/analyses/{newer_id}", owner_token).status_code, 404)
        self.assertTrue(shared_cache.is_file())

    def test_job_analysis_shortlist_restore_and_reports(self):
        owner_token = self.tokens[self.recruiter.id]
        created = self.request("POST", "/api/v1/jobs", owner_token, json={
            "title": "Platform Engineer", "department": "Engineering", "location": "Remote",
            "job_description": "Build reliable services and automate recruitment workflows for our engineering team.",
            "required_skills": ["Python", "SQL"],
        })
        self.assertEqual(created.status_code, 201, created.text)
        job = created.json()
        self.assertEqual(job["shortlist_threshold"], 75)
        saved_threshold = self.request("PATCH", f"/api/v1/jobs/{job['id']}/shortlist-threshold", owner_token, json={"threshold": 82})
        self.assertEqual(saved_threshold.status_code, 200, saved_threshold.text)
        self.assertEqual(saved_threshold.json()["threshold"], 82)
        self.assertEqual(self.request("PATCH", f"/api/v1/jobs/{job['id']}/shortlist-threshold", owner_token, json={"threshold": 101}).status_code, 422)
        self.assertEqual(self.request("GET", "/api/v1/jobs", owner_token).json()[0]["id"], job["id"])
        self.assertEqual(self.request("GET", "/api/v1/jobs", self.tokens[self.admin.id]).json()[0]["id"], job["id"])
        self.assertEqual(self.request("GET", "/api/v1/jobs", self.tokens[self.other.id]).json(), [])
        self.assertEqual(self.request("GET", f"/api/v1/jobs/{job['id']}", self.tokens[self.other.id]).status_code, 404)
        self.assertEqual(self.request("GET", "/api/v1/jobs").status_code, 401)
        replacement_description = "Updated role requirements from the analysis page PDF upload for this engineering position."
        updated_job = self.request("PATCH", f"/api/v1/jobs/{job['id']}", owner_token, json={"job_description":replacement_description})
        self.assertEqual(updated_job.status_code, 200, updated_job.text)
        job = updated_job.json()
        self.assertEqual(job["job_description"], replacement_description)

        upload = {"resumes": ("resume.pdf", BytesIO(b"%PDF-1.4 test resume"), "application/pdf")}
        response = self.request("POST", "/api/v1/analyses", owner_token, data={"job_id": job["id"]}, files=upload)
        self.assertEqual(response.status_code, 202, response.text)
        analysis_id = response.json()["id"]
        self.assertEqual(response.json()["run_number"], 1)
        self.assertEqual(self.request("GET", f"/api/v1/jobs/{job['id']}/analyses", self.tokens[self.admin.id]).json()[0]["id"], analysis_id)
        analysis = self.db.get(Analysis, analysis_id)
        self.assertEqual(analysis.job_id, job["id"])
        self.assertEqual(analysis.job_snapshot["description"], job["job_description"])
        self.assertEqual(analysis.job_snapshot["shortlist_threshold"], 82)
        original_created_at = analysis.created_at
        newer_batch = Analysis(job_id=job["id"], owner_user_id=self.recruiter.id, run_number=2, job_title=job["title"], job_description=job["job_description"], requirements={}, status="completed", created_at=original_created_at + timedelta(seconds=1))
        self.db.add(newer_batch); self.db.commit()
        ordered = self.request("GET", f"/api/v1/jobs/{job['id']}/analyses", self.tokens[self.admin.id]).json()
        self.assertEqual([item["id"] for item in ordered[:2]], [newer_batch.id, analysis_id])
        self.assertEqual([item["run_number"] for item in ordered[:2]], [2, 1])
        self.assertEqual(datetime.fromisoformat(ordered[1]["created_at"]).replace(tzinfo=None), original_created_at)
        self.assertEqual(self.request("GET", f"/api/v1/analyses/{analysis_id}", self.tokens[self.other.id]).status_code, 404)

        candidate = Candidate(analysis_id=analysis_id, filename="resume.pdf", name="=Formula Candidate", email="candidate@example.test", text="", skills=["Python"], missing_skills=[], entities={}, projects=[], education="BS", experience_details=[], semantic_score=80, keyword_score=70, experience_score=60, overall_score=75, recommendation="Good Match", insight="", is_shortlisted=False)
        unselected = Candidate(analysis_id=analysis_id, filename="other.pdf", name="Unselected Candidate", email="other-candidate@example.test", text="", skills=["SQL"], missing_skills=[], entities={}, projects=[], education="MS", experience_details=[], semantic_score=40, keyword_score=40, experience_score=40, overall_score=40, recommendation="Weak Match", insight="", is_shortlisted=False)
        no_email = Candidate(analysis_id=analysis_id, filename="no-email.pdf", name="No Email", email=None, text="", skills=[], missing_skills=[], entities={}, projects=[], education="", experience_details=[], semantic_score=20, keyword_score=20, experience_score=20, overall_score=20, recommendation="Weak Match", insight="", is_shortlisted=False)
        self.db.add_all([candidate, unselected, no_email]); self.db.commit(); candidate_id = candidate.id
        analysis.status = "completed"; self.db.commit()
        self.assertEqual(self.request("POST", f"/api/v1/candidates/{candidate_id}/shortlist", owner_token, json={"is_shortlisted": True}).status_code, 200)
        self.db.expire_all()
        self.assertTrue(self.db.get(Candidate, candidate_id).is_shortlisted)
        restored = self.request("GET", f"/api/v1/analyses/{analysis_id}", owner_token).json()
        self.assertTrue(restored["candidates"][0]["is_shortlisted"])
        self.assertEqual(restored["candidates"][0]["shortlist_source"], "manual")
        original_sender = main.EmailService.send_emails
        sent_payloads = []
        main.EmailService.send_emails = staticmethod(lambda payloads: (sent_payloads.extend(payloads) or {"success": True, "sent_count": len(payloads), "failed_count": 0, "errors": []}))
        try:
            email_response = self.request("POST", "/api/v1/candidates/send-email", owner_token, json={"recipient_ids":[candidate_id],"template_type":"custom","custom_subject":"Interview","custom_body":"Please reply."})
        finally:
            main.EmailService.send_emails = staticmethod(original_sender)
        self.assertEqual(email_response.status_code, 200, email_response.text)
        self.assertEqual([item["candidate_id"] for item in sent_payloads], [candidate_id])
        self.assertTrue(self.request("GET", f"/api/v1/analyses/{analysis_id}", owner_token).json()["candidates"][0]["email_sent"])
        main.EmailService.send_emails = staticmethod(lambda payloads: {"success": True, "sent_count": len(payloads), "failed_count": 0, "errors": []})
        try:
            mixed_email = self.request("POST", "/api/v1/candidates/send-email", owner_token, json={"recipient_ids":[candidate_id, no_email.id],"template_type":"custom","custom_subject":"Interview","custom_body":"Please reply."})
        finally:
            main.EmailService.send_emails = staticmethod(original_sender)
        self.assertEqual(mixed_email.status_code, 200, mixed_email.text)
        self.assertEqual(mixed_email.json()["sent_count"], 1)
        self.assertEqual(mixed_email.json()["skipped_count"], 1)
        original_sender = main.EmailService.send_emails
        main.EmailService.send_emails = staticmethod(lambda _payloads: {"success": False, "message": "Email service not configured", "sent_count": 0, "failed_count": 1, "errors": []})
        try:
            failed_email = self.request("POST", "/api/v1/candidates/send-email", owner_token, json={"recipient_ids":[unselected.id],"template_type":"custom","custom_subject":"Interview","custom_body":"Please reply."})
        finally:
            main.EmailService.send_emails = staticmethod(original_sender)
        self.assertEqual(failed_email.json()["failed_count"], 1)
        self.assertFalse(next(item for item in self.request("GET", f"/api/v1/analyses/{analysis_id}", owner_token).json()["candidates"] if item["id"] == unselected.id)["email_sent"])
        self.assertEqual(self.request("GET", f"/api/v1/analyses/{analysis_id}", self.tokens[self.admin.id]).json()["candidates"][0]["shortlist_source"], "manual")
        self.request("PATCH", f"/api/v1/jobs/{job['id']}/shortlist-threshold", owner_token, json={"threshold": 100})
        self.assertTrue(self.request("GET", f"/api/v1/analyses/{analysis_id}", owner_token).json()["candidates"][0]["is_shortlisted"])
        selected_export = self.request("POST", f"/api/v1/analyses/{analysis_id}/candidates/export.xlsx", owner_token, json={"candidate_ids":[candidate_id]})
        self.assertEqual(selected_export.status_code, 200, selected_export.text)
        selected_sheet = load_workbook(BytesIO(selected_export.content)).active
        self.assertEqual(selected_sheet.max_row, 2)
        self.assertEqual(selected_sheet.cell(2, 1).value, candidate_id)
        self.assertTrue(selected_sheet.cell(2, 13).hyperlink.target.endswith(f"/dashboard/resume/{candidate_id}"))
        resume_response = self.request("GET", f"/api/v1/candidates/{candidate_id}/resume", owner_token)
        self.assertEqual(resume_response.status_code, 200)
        self.assertEqual(resume_response.headers["content-type"], "application/pdf")
        self.assertEqual(self.request("GET", f"/api/v1/candidates/{candidate_id}/resume").status_code, 401)
        self.assertEqual(self.request("POST", f"/api/v1/analyses/{analysis_id}/candidates/export.xlsx", owner_token, json={"candidate_ids":["unknown"]}).status_code, 404)
        full = self.request("GET", f"/api/v1/analyses/{analysis_id}/report.xlsx", owner_token)
        self.assertEqual(full.status_code, 200)
        self.assertEqual(self.request("GET", f"/api/v1/analyses/{analysis_id}/report.xlsx").status_code, 401)
        self.assertEqual(self.request("GET", "/api/v1/analyses/does-not-exist/report.xlsx", owner_token).status_code, 404)
        sheet = load_workbook(BytesIO(full.content), data_only=False).get_sheet_by_name("Summary")
        self.assertIn(candidate_id, [cell.value for row in sheet.iter_rows() for cell in row])
        self.assertNotEqual(next(cell.value for row in sheet.iter_rows() for cell in row if cell.value == "'=Formula Candidate"), "=Formula Candidate")
        shortlist = self.request("GET", f"/api/v1/analyses/{analysis_id}/shortlisted-report.xlsx", owner_token)
        self.assertEqual(shortlist.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(shortlist.content)).active.max_row, 2)
        self.assertEqual(self.request("POST", f"/api/v1/analyses/{analysis_id}/shortlist/bulk-remove", owner_token, json={"candidate_ids":[candidate_id]}).json()["updated_count"], 1)
        self.assertFalse(self.request("GET", f"/api/v1/analyses/{analysis_id}", owner_token).json()["candidates"][0]["is_shortlisted"])
        bulk_added = self.request("POST", f"/api/v1/analyses/{analysis_id}/shortlist/bulk-add", owner_token, json={"candidate_ids":[candidate_id,candidate_id]})
        self.assertEqual(bulk_added.status_code, 200, bulk_added.text)
        self.assertEqual(bulk_added.json()["updated_count"], 1)
        restored_manual = self.request("GET", f"/api/v1/analyses/{analysis_id}", owner_token).json()["candidates"]
        self.assertTrue(next(item for item in restored_manual if item["id"] == candidate_id)["is_shortlisted"])
        self.assertEqual(next(item for item in restored_manual if item["id"] == candidate_id)["shortlist_source"], "manual")
        self.assertEqual(self.request("POST", f"/api/v1/analyses/{analysis_id}/shortlist/bulk-add", self.tokens[self.other.id], json={"candidate_ids":[candidate_id]}).status_code, 404)
        self.request("PATCH", f"/api/v1/jobs/{job['id']}/shortlist-threshold", owner_token, json={"threshold": 0})
        latest_candidates = self.request("GET", f"/api/v1/analyses/{analysis_id}", owner_token).json()["candidates"]
        self.assertTrue(next(item for item in latest_candidates if item["id"] == candidate_id)["is_shortlisted"])
        self.assertEqual(self.request("GET", f"/api/v1/analyses/{analysis_id}", self.tokens[self.admin.id]).json()["candidates"][0]["is_shortlisted"], True)
        self.assertEqual(self.request("GET", "/api/v1/jobs", self.tokens[self.admin.id]).json()[0]["shortlist_threshold"], 0)
        self.assertEqual(self.request("POST", f"/api/v1/analyses/{analysis_id}/candidates/export.xlsx", self.tokens[self.other.id], json={"candidate_ids":[candidate_id]}).status_code, 404)
        self.assertEqual(self.request("POST", f"/api/v1/candidates/{candidate_id}/shortlist", owner_token, json={"is_shortlisted": False}).status_code, 200)
        self.assertEqual(self.request("GET", f"/api/v1/analyses/{analysis_id}/shortlisted-report.xlsx", owner_token).status_code, 200)
        self.assertEqual(self.request("PATCH", f"/api/v1/jobs/{job['id']}", owner_token, json={"status":"closed"}).status_code, 200)
        self.assertEqual(self.request("POST", "/api/v1/analyses", owner_token, data={"job_id":job["id"]}, files=upload).status_code, 409)
        reopened=self.request("PATCH", f"/api/v1/jobs/{job['id']}", owner_token, json={"status":"open"})
        self.assertEqual(reopened.status_code,200)
        self.assertEqual(self.request("GET", "/api/v1/jobs", self.tokens[self.admin.id]).json()[0]["status"],"open")

    def test_legacy_upgrade_is_additive_and_idempotent(self):
        legacy = create_engine(f"sqlite:///{Path(_temp.name) / 'legacy.db'}")
        with legacy.begin() as connection:
            connection.execute(text("CREATE TABLE analyses (id VARCHAR(36) PRIMARY KEY, job_id VARCHAR(36), created_at TIMESTAMP)"))
            connection.execute(text("CREATE TABLE candidates (id VARCHAR(36) PRIMARY KEY)"))
            connection.execute(text("CREATE TABLE users (id VARCHAR(36) PRIMARY KEY)"))
            connection.execute(text("CREATE TABLE jobs (id VARCHAR(36) PRIMARY KEY)"))
            connection.execute(text("INSERT INTO analyses (id, job_id, created_at) VALUES ('legacy-analysis', 'job-1', '2026-01-02 00:00:00')"))
            connection.execute(text("INSERT INTO analyses (id, job_id, created_at) VALUES ('legacy-older', 'job-1', '2026-01-01 00:00:00')"))
        from app.migrations import run_migrations
        run_migrations(legacy); run_migrations(legacy)
        self.assertIn("job_id", {col["name"] for col in inspect(legacy).get_columns("analyses")})
        self.assertIn("shortlist_source", {col["name"] for col in inspect(legacy).get_columns("candidates")})
        self.assertIn("shortlist_threshold", {col["name"] for col in inspect(legacy).get_columns("jobs")})
        self.assertIn("run_number", {col["name"] for col in inspect(legacy).get_columns("analyses")})
        self.assertIn("schema_migrations", inspect(legacy).get_table_names())
        with legacy.connect() as connection:
            self.assertEqual(connection.execute(text("SELECT id FROM analyses WHERE id = 'legacy-analysis'")).scalar_one(), "legacy-analysis")
            self.assertEqual(connection.execute(text("SELECT run_number FROM analyses WHERE id = 'legacy-analysis'")).scalar_one(), 2)
            self.assertEqual(connection.execute(text("SELECT run_number FROM analyses WHERE id = 'legacy-older'")).scalar_one(), 1)
        legacy.dispose()


if __name__ == "__main__":
    unittest.main()
